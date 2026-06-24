"""Sumber data: Excel upload, dan Google Sheet.

Keduanya mengembalikan list[dict] dengan key = header kolom (harus cocok
dengan 'key' field di templates.yaml).
"""
import openpyxl


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_excel(path: str, sheet: str | None = None) -> list[dict]:
    """Baca Excel. Baris pertama (yang ada isinya) = header, sisanya = data."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active

    rows = list(ws.iter_rows(values_only=True))
    # cari baris header pertama yang punya >=2 sel terisi
    header_idx = None
    for i, r in enumerate(rows):
        if sum(1 for c in r if c not in (None, "")) >= 2:
            header_idx = i
            break
    if header_idx is None:
        return []

    headers = [_norm(c) for c in rows[header_idx]]
    out = []
    for r in rows[header_idx + 1:]:
        if not any(c not in (None, "") for c in r):
            continue
        rec = {}
        for h, c in zip(headers, r):
            if h:
                rec[h] = _norm(c)
        out.append(rec)
    return out


def read_gsheet(sheet_id: str, credentials_path: str, worksheet: str | None = None) -> list[dict]:
    """Baca Google Sheet via service account. Baris pertama = header."""
    import gspread
    from google.oauth2.service_account import Credentials

    scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
    creds = Credentials.from_service_account_file(credentials_path, scopes=scopes)
    gc = gspread.authorize(creds)
    sh = gc.open_by_key(sheet_id)
    ws = sh.worksheet(worksheet) if worksheet else sh.sheet1
    records = ws.get_all_records()  # list[dict] pakai baris1 sebagai header
    return [{k: _norm(v) for k, v in rec.items()} for rec in records]


def validate_rows(rows: list[dict], required_keys: list[str]) -> tuple[bool, str]:
    """Cek apakah kolom yang dibutuhkan ada. Return (ok, pesan)."""
    if not rows:
        return False, "Tidak ada data yang terbaca."
    present = set(rows[0].keys())
    missing = [k for k in required_keys if k not in present]
    if missing:
        return False, "Kolom belum ada di data: " + ", ".join(missing)
    return True, f"{len(rows)} baris data terbaca."
